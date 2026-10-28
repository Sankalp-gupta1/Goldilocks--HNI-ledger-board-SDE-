"""Authenticated Flask API and the compiled React workspace."""

import csv
import io
import json
import os
import re
import secrets
import uuid
from collections import defaultdict
from datetime import date
from pathlib import Path
from flask import Flask, abort, g, jsonify, request, send_file, send_from_directory
from werkzeug.exceptions import HTTPException
from werkzeug.middleware.proxy_fix import ProxyFix
from werkzeug.security import generate_password_hash
from . import auth, accounting, ingestion, service

ROOT = Path(__file__).resolve().parents[1]


def create_app(database_path=None, testing=False):
    import hni_accounting_system as core
    import hni_ledger_extension as ext

    app = Flask(__name__, static_folder=None)
    app.config.update(
        TESTING=testing,
        MAX_CONTENT_LENGTH=16_000_000,
        MAX_FORM_MEMORY_SIZE=64_000,
        MAX_FORM_PARTS=20,
        COOKIE_SECURE=(
            os.getenv("APP_ENV") == "production" or os.getenv("RENDER") == "true"
        ),
    )
    if os.getenv("RENDER") == "true":
        app.wsgi_app = ProxyFix(app.wsgi_app, x_proto=1)
    db = core.DBStore(database_path or core.DB_PATH)
    ext.extend_db(db)
    edb = ext.ExtendedDBStore(db)
    auth.initialize(db)
    accounting.initialize(db)
    service.initialize(db)
    app.extensions["ledger"] = {"db": db, "edb": edb}
    # Migrate old approved rows to real double-entry journals without changing classifications.
    with db._conn() as c:
        c.execute("BEGIN IMMEDIATE")
        for row in c.execute(
            "SELECT * FROM ledger WHERE NOT EXISTS(SELECT 1 FROM journal_entries j WHERE j.ledger_id=ledger.id)"
        ).fetchall():
            if row["txn_date"] and row["amount"] > 0:
                accounting.post(c, dict(row))
    dummy = auth.password_hash("fictional-dummy-password", "fictional-dummy-password")

    @app.before_request
    def boundary():
        g.request_id = uuid.uuid4().hex
        if (
            request.is_json
            and request.content_length
            and request.content_length > 128_000
        ):
            abort(413)
        if request.method in ("POST", "PUT", "PATCH", "DELETE"):
            if request.headers.get("X-Goldilocks") != "1":
                abort(
                    403,
                    description="Use the Goldilocks workspace to make this request.",
                )
            origin = request.headers.get("Origin")
            if origin and origin.rstrip("/") != request.host_url.rstrip("/"):
                abort(403, description="The request origin is not allowed.")
        public = {
            "/api/health",
            "/api/session",
            "/api/register",
            "/api/login",
            "/api/recover",
        }
        if request.path.startswith("/api/") and request.path not in public:
            auth.require_user(db)

    @app.after_request
    def headers(response):
        response.headers["X-Content-Type-Options"] = "nosniff"
        response.headers["X-Frame-Options"] = "DENY"
        response.headers["Referrer-Policy"] = "no-referrer"
        response.headers["Permissions-Policy"] = (
            "camera=(),microphone=(),geolocation=()"
        )
        response.headers["Content-Security-Policy"] = (
            "default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; font-src 'self'; connect-src 'self'; object-src 'none'; frame-ancestors 'none'; base-uri 'self'; form-action 'self'"
        )
        if request.path.startswith("/api/") or request.path == "/":
            response.headers["Cache-Control"] = "private, no-store"
        if app.config["COOKIE_SECURE"]:
            response.headers["Strict-Transport-Security"] = "max-age=31536000"
        return response

    @app.errorhandler(Exception)
    def error(error):
        if isinstance(error, HTTPException):
            code = error.code
            message = error.description
        elif isinstance(error, LookupError):
            code = 404
            message = "This record is not available in your workspace."
        elif isinstance(error, ValueError):
            code = 422
            message = str(error)
        else:
            code = 500
            message = "This operation could not finish. Please retry."
            app.logger.error(
                "request_failed id=%s type=%s",
                getattr(g, "request_id", ""),
                type(error).__name__,
            )
        return jsonify(error=message, request_id=getattr(g, "request_id", "")), code

    def body():
        value = request.get_json(silent=True)
        if not isinstance(value, dict):
            abort(422, description="Send a JSON object.")
        return value

    def filters():
        start = request.args.get("from") or None
        end = request.args.get("to") or None
        for value in (start, end):
            if value:
                try:
                    date.fromisoformat(value)
                except ValueError:
                    abort(422, description="Use dates in YYYY-MM-DD format.")
        if start and end and start > end:
            abort(422, description="The start date must be before the end date.")
        return start, end, request.args.get("account") or None

    def ledger_rows():
        start, end, account = filters()
        where = "user_id=?"
        params = [g.user_id]
        for clause, value in [
            ("txn_date>=?", start),
            ("txn_date<=?", end),
            ("account_id=?", account),
        ]:
            if value:
                where += " AND " + clause
                params.append(value)
        search = request.args.get("q", "").strip()[:120]
        if search:
            where += (
                " AND (narration LIKE ? ESCAPE '\\' OR counterparty LIKE ? ESCAPE '\\')"
            )
            term = (
                "%"
                + search.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
                + "%"
            )
            params.extend([term, term])
        with db._conn() as c:
            return [
                dict(x)
                for x in c.execute(
                    f"SELECT * FROM ledger WHERE {where} ORDER BY txn_date DESC,created_at DESC LIMIT 20000",
                    params,
                )
            ]

    @app.get("/api/health")
    def health():
        import shutil

        with db._conn() as c:
            c.execute("SELECT 1").fetchone()
        ephemeral = os.getenv("HNI_STORAGE_MODE") == "ephemeral" or (
            os.getenv("RENDER") == "true" and not os.getenv("HNI_DATA_DIR")
        )
        return jsonify(
            storage_warning="This host uses temporary storage. Accounts and records can be lost on restart. Use fictional files until a persistent disk is configured."
            if ephemeral
            else None,
            status="ok",
            version="2.0.0",
            storage="sqlite",
            ocr_available=bool(shutil.which("tesseract")),
            persistent_storage_configured=bool(os.getenv("HNI_DATA_DIR"))
            and not ephemeral,
            external_financial_apis=False,
        )

    @app.get("/api/session")
    def session():
        user = auth.current_user(db)
        return jsonify(user=auth.public_user(user) if user else None)

    @app.post("/api/register")
    def register():
        p = body()
        email = str(p.get("email", "")).strip().lower()
        name = str(p.get("name", "")).strip()
        if (
            not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", email)
            or len(email) > 254
            or not 2 <= len(name) <= 100
        ):
            abort(422, description="Enter your name and a valid email address.")
        auth.limit(db, "register:" + email, 6)
        auth.limit(db, "register:global", 100)
        hashed = auth.password_hash(p.get("password"), p.get("confirm_password"))
        uid = uuid.uuid4().hex
        with db._conn() as c:
            c.execute("BEGIN IMMEDIATE")
            if c.execute(
                "SELECT 1 FROM users WHERE lower(email)=?", (email,)
            ).fetchone():
                abort(409, description="Could not create this account. Try signing in.")
            c.execute(
                "INSERT INTO users(id,name,email,phone,created_at,user_type,password_hash) VALUES(?,?,?,?,?,?,?)",
                (uid, name, email, "", date.today().isoformat(), "INDIVIDUAL", hashed),
            )
        edb.ensure_user_profile_exists(uid)
        edb.upsert_user_profile(uid, {"entity_type": "INDIVIDUAL", "legal_name": name})
        code = auth.recovery(db, uid)
        response = jsonify(user=auth.public_user(edb.get_user(uid)), recovery_code=code)
        response.status_code = 201
        return auth.new_session(db, uid, response)

    @app.post("/api/login")
    def login():
        p = body()
        identifier = str(p.get("email", "")).strip().lower()
        password = str(p.get("password", ""))
        auth.limit(db, "login:" + identifier, 15)
        auth.limit(db, "login:global", 400)
        with db._conn() as c:
            c.execute("BEGIN IMMEDIATE")
            row = c.execute(
                "SELECT * FROM users WHERE lower(email)=? OR phone=? LIMIT 1",
                (identifier, identifier),
            ).fetchone()
            if not auth.verify(password, row["password_hash"] if row else dummy):
                abort(401, description="Email or password is incorrect.")
            user = dict(row)
            if user["password_hash"].startswith("pbkdf2_sha256$"):
                c.execute(
                    "UPDATE users SET password_hash=? WHERE id=?",
                    (generate_password_hash(password, method="scrypt"), user["id"]),
                )
            token = auth.issue_session(c, user["id"])
        return auth.session_cookie(token, jsonify(user=auth.public_user(user)))

    @app.post("/api/recover")
    def recover():
        import hmac

        p = body()
        email = str(p.get("email", "")).strip().lower()
        auth.limit(db, "recover:" + email, 8)
        auth.limit(db, "recover:global", 100)
        hashed = auth.password_hash(p.get("password"), p.get("confirm_password"))
        code = secrets.token_urlsafe(32)
        with db._conn() as c:
            c.execute("BEGIN IMMEDIATE")
            row = c.execute(
                "SELECT u.id,r.code_hash FROM users u JOIN web_recovery r ON r.user_id=u.id WHERE lower(u.email)=?",
                (email,),
            ).fetchone()
            if not row or not hmac.compare_digest(
                row["code_hash"], auth.digest(str(p.get("recovery_code", "")))
            ):
                abort(403, description="The recovery details could not be verified.")
            c.execute(
                "UPDATE users SET password_hash=?,reset_token=NULL,reset_token_expires_at=NULL WHERE id=?",
                (hashed, row["id"]),
            )
            c.execute("DELETE FROM web_sessions WHERE user_id=?", (row["id"],))
            c.execute(
                "UPDATE web_recovery SET code_hash=? WHERE user_id=?",
                (auth.digest(code), row["id"]),
            )
            token = auth.issue_session(c, row["id"])
        return auth.session_cookie(
            token,
            jsonify(user=auth.public_user(edb.get_user(row["id"])), recovery_code=code),
        )

    @app.post("/api/logout")
    def logout():
        with db._conn() as c:
            c.execute(
                "DELETE FROM web_sessions WHERE token_hash=?",
                (auth.digest(request.cookies.get(auth.COOKIE, "")),),
            )
        response = jsonify(ok=True)
        response.delete_cookie(auth.COOKIE, path="/")
        return response

    @app.get("/api/meta")
    def meta():
        ml = core.get_ml()
        return jsonify(
            categories=[
                dict(key=k, book=v[0], section=v[1], group=v[2], label=v[3])
                for k, v in core.LEDGER_MAP.items()
                if k not in ("exp_loan_emi", "exp_credit_card")
            ],
            banks=list(ingestion.BANKS),
            model={
                "ready": ml.ready,
                "classes": len(ml.labels),
                "training": "Authored synthetic examples; not real customer data.",
                "score_note": "A model score is not verified accuracy.",
            },
        )

    @app.get("/api/overview")
    def overview():
        rows = ledger_rows()
        start, end, account = filters()
        report = accounting.report(db, g.user_id, start, end, account)
        months = defaultdict(
            lambda: {"income": 0, "expense": 0, "inflow": 0, "outflow": 0}
        )
        for r in reversed(rows):
            m = months[r["txn_date"][:7]]
            amount = accounting.paise(r["amount"]) / 100
            m["inflow" if r["txn_type"] == "credit" else "outflow"] += amount
        for month, values in report["monthly_performance"].items():
            months[month].update(values)
        with db._conn() as c:
            sources = c.execute(
                "SELECT count(*) FROM raw_import_batches WHERE user_id=?", (g.user_id,)
            ).fetchone()[0]
            accounts = [
                dict(x)
                for x in c.execute(
                    """SELECT account_id,count(*) AS transactions FROM ledger WHERE user_id=? GROUP BY account_id UNION SELECT account_id,0 FROM pending_ledger p WHERE user_id=? AND NOT EXISTS(SELECT 1 FROM ledger l WHERE l.account_id=p.account_id AND l.user_id=p.user_id) GROUP BY account_id""",
                    (g.user_id, g.user_id),
                )
            ]
            events = [
                dict(x)
                for x in c.execute(
                    "SELECT id,action,record_id,created_at FROM review_events WHERE user_id=? ORDER BY created_at DESC,rowid DESC LIMIT 8",
                    (g.user_id,),
                )
            ]
        return jsonify(
            report=report,
            months=[
                dict(month=m, **{k: round(v, 2) for k, v in values.items()})
                for m, values in sorted(months.items())
            ],
            recent=rows[:6],
            transactions=len(rows),
            statements=sources,
            accounts=accounts,
            events=events,
            reconciliation=service.reconcile(db, g.user_id),
        )

    @app.get("/api/pending")
    def pending():
        return jsonify(rows=edb.get_pending_transactions(g.user_id))

    @app.get("/api/ledger")
    def ledger():
        return jsonify(rows=ledger_rows())

    @app.post("/api/approve")
    def approve():
        return jsonify(approved=service.approve(db, g.user_id, body().get("ids", [])))

    @app.post("/api/transactions/<rid>/classification")
    def correct(rid):
        p = body()
        return jsonify(
            row=service.reclassify(
                db,
                g.user_id,
                rid,
                str(p.get("key", "")),
                str(p.get("note", ""))[:1000],
                bool(p.get("posted")),
            )
        )

    @app.post("/api/import")
    def upload():
        auth.limit(db, "upload:" + g.user_id, 40)
        f = request.files.get("file")
        if not f:
            abort(422, description="Choose a statement to import.")
        name = Path((f.filename or "statement").replace("\\", "/")).name[:180]
        data = f.read()
        parsed = ingestion.parse_statement(data, name, request.form.get("password", ""))
        for k in ("opening_balance", "closing_balance"):
            if request.form.get(k, "").strip():
                parsed[k] = ingestion.money(request.form[k])
                parsed[k + "_source"] = "user"
        return jsonify(
            service.stage(
                db,
                edb,
                g.user_id,
                request.form.get("account_id", "Primary account").strip(),
                name,
                data,
                parsed,
            )
        ), 201

    @app.post("/api/sample")
    def sample():
        data = (ROOT / "samples" / "fictional-statement.csv").read_bytes()
        parsed = ingestion.parse_statement(data, "fictional-statement.csv")
        parsed["opening_balance"] = 100000
        parsed["opening_balance_source"] = "fixture declaration"
        parsed["warnings"] = [
            w
            for w in parsed["warnings"]
            if not w.startswith("Opening balance is inferred")
        ]
        return jsonify(
            service.stage(
                db,
                edb,
                g.user_id,
                "Sample · HDFC".replace(" · ", " - "),
                "fictional-statement.csv",
                data,
                parsed,
            )
        ), 201

    @app.get("/api/statement-template")
    def template():
        return send_file(
            ROOT / "samples" / "statement-template.csv",
            as_attachment=True,
            download_name="statement-template.csv",
        )

    @app.get("/api/statements")
    def statements():
        with db._conn() as c:
            rows = [
                dict(x)
                for x in c.execute(
                    """SELECT b.*,s.bank,s.parser,s.warnings FROM raw_import_batches b LEFT JOIN source_documents s ON s.batch_id=b.id WHERE b.user_id=? ORDER BY b.created_at DESC""",
                    (g.user_id,),
                )
            ]
        return jsonify(rows=rows)

    @app.get("/api/statements/<rid>/source")
    def original(rid):
        with db._conn() as c:
            row = c.execute(
                "SELECT s.content,b.source_file_name FROM source_documents s JOIN raw_import_batches b ON b.id=s.batch_id WHERE s.batch_id=? AND s.user_id=?",
                (rid, g.user_id),
            ).fetchone()
        if not row:
            abort(404)
        return send_file(
            io.BytesIO(row["content"]),
            mimetype="application/octet-stream",
            as_attachment=True,
            download_name=row["source_file_name"],
        )

    @app.get("/api/trace/<rid>")
    def trace(rid):
        with db._conn() as c:
            row = c.execute(
                "SELECT * FROM ledger WHERE id=? AND user_id=?", (rid, g.user_id)
            ).fetchone()
            posted = bool(row)
            if not row:
                row = c.execute(
                    "SELECT * FROM pending_ledger WHERE id=? AND user_id=?",
                    (rid, g.user_id),
                ).fetchone()
            if not row:
                abort(404)
            raw = c.execute(
                "SELECT * FROM raw_import_rows WHERE moved_to_pending_id=? AND user_id=?",
                (rid, g.user_id),
            ).fetchone()
            journal = [
                dict(x)
                for x in c.execute(
                    "SELECT l.* FROM journal_lines l JOIN journal_entries j ON j.id=l.entry_id WHERE j.ledger_id=? AND j.user_id=?",
                    (rid, g.user_id),
                )
            ]
            events = [
                dict(x)
                for x in c.execute(
                    "SELECT * FROM review_events WHERE record_id=? AND user_id=? ORDER BY created_at,rowid",
                    (rid, g.user_id),
                )
            ]
            splits = [
                dict(x)
                for x in c.execute(
                    "SELECT * FROM transaction_splits WHERE transaction_id=? AND user_id=?",
                    (rid, g.user_id),
                )
            ]
        return jsonify(
            splits=splits,
            transaction=dict(row),
            posted=posted,
            source=dict(raw) if raw else None,
            journal=journal,
            events=events,
        )

    @app.post("/api/transactions/<rid>/split")
    def split(rid):
        p = body()
        return jsonify(
            service.split_transaction(
                db,
                g.user_id,
                rid,
                p.get("allocations"),
                str(p.get("note", ""))[:1000],
                bool(p.get("posted")),
            )
        )

    @app.get("/api/openings")
    def openings():
        with db._conn() as c:
            bank = [
                dict(x)
                for x in c.execute(
                    "SELECT * FROM opening_positions WHERE user_id=?", (g.user_id,)
                )
            ]
            holdings = [
                dict(x)
                for x in c.execute(
                    "SELECT * FROM opening_holdings WHERE user_id=?", (g.user_id,)
                )
            ]
        return jsonify(bank=bank, holdings=holdings)

    @app.put("/api/openings")
    def set_opening():
        p = body()
        account = str(p.get("account_id", "")).strip()
        key = p.get("key", "bank")
        day = str(p.get("as_of", ""))
        if not re.fullmatch(r"[A-Za-z0-9 _.-]{1,60}", account):
            raise ValueError("Enter a valid account label.")
        date.fromisoformat(day)
        value = accounting.paise(p.get("amount"))
        if key != "bank" and (
            key not in core.LEDGER_MAP
            or core.LEDGER_MAP[key][1] not in ("Assets", "Liabilities")
        ):
            raise ValueError("Use an asset or liability category.")
        if key != "bank" and value < 0:
            raise ValueError("Enter a positive opening holding or debt.")
        with db._conn() as c:
            c.execute("BEGIN IMMEDIATE")
            earliest = c.execute(
                "SELECT min(txn_date) FROM ledger WHERE user_id=? AND account_id=?",
                (g.user_id, account),
            ).fetchone()[0]
            if earliest and day > earliest:
                raise ValueError(
                    "An opening must be dated on or before the first posted transaction in this account."
                )
            if key == "bank":
                c.execute(
                    "INSERT INTO opening_positions VALUES(?,?,?,?) ON CONFLICT(user_id,account_id) DO UPDATE SET amount=excluded.amount,as_of=excluded.as_of",
                    (g.user_id, account, value, day),
                )
            else:
                c.execute(
                    "INSERT INTO opening_holdings VALUES(?,?,?,?,?) ON CONFLICT(user_id,account_id,ledger_key) DO UPDATE SET amount=excluded.amount,as_of=excluded.as_of",
                    (g.user_id, account, key, value, day),
                )
            service.event(
                c,
                g.user_id,
                "opening_position_saved",
                account,
                after={"key": key, "amount": value / 100, "as_of": day},
            )
        return jsonify(ok=True)

    @app.put("/api/statements/<rid>/balances")
    def confirm_balances(rid):
        p = body()
        opening = ingestion.money(p.get("opening"))
        closing = ingestion.money(p.get("closing"))
        if opening is None or closing is None:
            raise ValueError("Enter both balances from the statement summary.")
        with db._conn() as c:
            c.execute("BEGIN IMMEDIATE")
            before = c.execute(
                "SELECT * FROM raw_import_batches WHERE id=? AND user_id=?",
                (rid, g.user_id),
            ).fetchone()
            if not before:
                abort(404)
            c.execute(
                "UPDATE raw_import_batches SET opening_balance=?,closing_balance=? WHERE id=? AND user_id=?",
                (opening, closing, rid, g.user_id),
            )
            c.execute(
                "UPDATE source_documents SET opening_source='user confirmed',closing_source='user confirmed' WHERE batch_id=? AND user_id=?",
                (rid, g.user_id),
            )
            c.execute(
                "UPDATE opening_positions SET amount=? WHERE user_id=? AND account_id=? AND as_of=?",
                (
                    accounting.paise(opening),
                    g.user_id,
                    before["account_id"],
                    before["statement_from_date"],
                ),
            )
            service.event(
                c,
                g.user_id,
                "statement_balances_confirmed",
                rid,
                after={"opening": opening, "closing": closing},
            )
        return jsonify(ok=True)

    @app.get("/api/reconciliation")
    def reconciliation():
        return jsonify(
            rows=service.reconcile(db, g.user_id),
            journal=accounting.report(db, g.user_id, *filters()),
        )

    @app.get("/api/reports")
    def reports():
        return jsonify(accounting.report(db, g.user_id, *filters()))

    @app.get("/api/export/<kind>")
    def export(kind):
        rows = ledger_rows()
        data = accounting.report(db, g.user_id, *filters())
        columns = [
            "txn_date",
            "account_id",
            "narration",
            "txn_type",
            "amount",
            "account",
            "counterparty",
            "source",
            "id",
        ]
        if kind == "csv":
            out = io.StringIO()
            writer = csv.DictWriter(out, fieldnames=columns)
            writer.writeheader()
            for row in rows:
                writer.writerow(
                    {
                        k: (
                            "'" + str(row.get(k, ""))
                            if isinstance(row.get(k), str)
                            and str(row[k]).startswith(("=", "+", "-", "@", "\t", "\r"))
                            else row.get(k, "")
                        )
                        for k in columns
                    }
                )
            return send_file(
                io.BytesIO(out.getvalue().encode("utf-8-sig")),
                mimetype="text/csv",
                as_attachment=True,
                download_name="goldilocks-ledger.csv",
            )
        if kind == "xlsx":
            from openpyxl import Workbook

            wb = Workbook()
            ws = wb.active
            ws.title = "Posted ledger"
            ws.append(columns)
            for row in rows:
                ws.append([row.get(k, "") for k in columns])
                for cell in ws[ws.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            out = io.BytesIO()
            wb.save(out)
            out.seek(0)
            return send_file(
                out,
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                as_attachment=True,
                download_name="goldilocks-ledger.xlsx",
            )
        if kind == "pdf":
            from reportlab.lib.pagesizes import A4
            from reportlab.platypus import (
                SimpleDocTemplate,
                Paragraph,
                Spacer,
                Table,
                TableStyle,
            )
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib import colors
            from xml.sax.saxutils import escape

            out = io.BytesIO()
            styles = getSampleStyleSheet()
            parts = [
                Paragraph("Goldilocks | Financial statements", styles["Title"]),
                Paragraph(escape(data["scope"]), styles["Normal"]),
                Spacer(1, 20),
            ]
            for title, entries in [
                ("Income", data["income"]),
                ("Expenditure", data["expenses"]),
                *data["sections"].items(),
            ]:
                parts.append(Paragraph(title, styles["Heading2"]))
                table = Table(
                    [
                        [
                            Paragraph(escape(x["label"]), styles["Normal"]),
                            f"INR {x['amount']:,.2f}",
                        ]
                        for x in entries
                    ]
                    or [["No recorded entries", "INR 0.00"]],
                    colWidths=[330, 140],
                )
                table.setStyle(
                    TableStyle(
                        [
                            ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                            ("LINEBELOW", (0, -1), (-1, -1), 0.4, colors.lightgrey),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ]
                    )
                )
                parts.extend([table, Spacer(1, 10)])
            parts.append(
                Paragraph(
                    f"Assets less liabilities and equity: INR {data['difference']:,.2f}. This is not an independent audit.",
                    styles["Normal"],
                )
            )
            SimpleDocTemplate(out, pagesize=A4).build(parts)
            out.seek(0)
            return send_file(
                out,
                mimetype="application/pdf",
                as_attachment=True,
                download_name="goldilocks-statements.pdf",
            )
        abort(404)

    @app.get("/api/rules")
    def rules():
        return jsonify(rows=edb.get_custom_rules(g.user_id))

    @app.post("/api/rules")
    def add_rule():
        p = body()
        pattern = str(p.get("pattern", "")).strip()
        mode = p.get("match_mode", "contains")
        if not 2 <= len(pattern) <= 160 or mode not in ("exact", "contains"):
            abort(
                422, description="Use an exact or contains rule with 2–160 characters."
            )
        if p.get("ledger_key") not in core.LEDGER_MAP or p.get("txn_type", "") not in (
            "",
            "debit",
            "credit",
        ):
            abort(422, description="Choose a valid category and direction.")
        if p.get("ledger_key") in ("exp_loan_emi", "exp_credit_card"):
            abort(
                422,
                description="EMIs need a principal/interest split. Card settlements must reduce card liability.",
            )
        row = edb.add_custom_rule(
            g.user_id,
            pattern,
            str(p.get("ledger_key", "")),
            str(p.get("txn_type", "")),
            mode,
        )
        with db._conn() as c:
            service.event(c, g.user_id, "rule_created", row["id"], after=row)
        return jsonify(row=row), 201

    @app.delete("/api/rules/<rid>")
    def delete_rule(rid):
        with db._conn() as c:
            if not c.execute(
                "UPDATE custom_classifier_rules SET is_active=0 WHERE id=? AND user_id=?",
                (rid, g.user_id),
            ).rowcount:
                abort(404)
            service.event(c, g.user_id, "rule_disabled", rid)
        return jsonify(ok=True)

    @app.get("/api/learning-examples")
    def learning_examples():
        with db._conn() as c:
            rows = [
                dict(x)
                for x in c.execute(
                    "SELECT narration,txn_type,ledger_key,reviewed_at FROM learning_examples WHERE user_id=?",
                    (g.user_id,),
                )
            ]
        return send_file(
            io.BytesIO(json.dumps(rows, indent=2).encode()),
            mimetype="application/json",
            as_attachment=True,
            download_name="my-reviewed-examples.json",
        )

    @app.get("/api/profile")
    def profile():
        return jsonify(
            profile=edb.get_user_profile(g.user_id),
            counterparties=edb.get_known_counterparties(g.user_id),
            accounts=edb.get_known_accounts(g.user_id),
        )

    @app.put("/api/profile")
    def save_profile():
        p = body()
        profile = edb.upsert_user_profile(g.user_id, p)
        return jsonify(profile=profile)

    @app.get("/api/trading")
    def trading():
        rows = [
            r
            for r in ledger_rows()
            if r["ledger_key"]
            in (
                "trading_funds_added",
                "trading_payout",
                "asset_broker_payout",
                "asset_broker_balance",
            )
        ]
        deposits = (
            sum(accounting.paise(r["amount"]) for r in rows if r["txn_type"] == "debit")
            / 100
        )
        payouts = (
            sum(
                accounting.paise(r["amount"]) for r in rows if r["txn_type"] == "credit"
            )
            / 100
        )
        return jsonify(
            rows=rows,
            deposits=deposits,
            payouts=payouts,
            net=round(deposits - payouts, 2),
            scope="Net broker cash transfers in the selected period, not investment profit or market portfolio value.",
        )

    @app.get("/")
    def index():
        folder = ROOT / "web"
        if not (folder / "index.html").exists():
            return "Build the frontend with npm run build --prefix frontend.", 503
        return send_from_directory(folder, "index.html")

    @app.get("/assets/<path:name>")
    def asset(name):
        return send_from_directory(ROOT / "web" / "assets", name, max_age=31536000)

    return app
