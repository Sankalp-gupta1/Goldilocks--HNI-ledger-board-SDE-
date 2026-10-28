"""Bounded statement ingestion with retained source rows and explicit ambiguity."""

import csv
import io
import re
import shutil
import zipfile
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

BANKS = {
    "HDFC": "hdfc",
    "ICICI": "icici",
    "SBI": "state bank|\\bsbi\\b",
    "Axis": "axis",
    "PNB": "punjab national|\\bpnb\\b",
    "Union": "union bank",
    "AU": "au (small|bank)",
    "IDFC First": "idfc",
}
ALIASES = {
    "date": [
        "date",
        "transaction date",
        "txn date",
        "tran date",
        "value date",
        "posted date",
    ],
    "description": [
        "description",
        "narration",
        "particulars",
        "transaction particulars",
        "remarks",
        "transaction remarks",
        "details",
    ],
    "debit": [
        "debit",
        "debit amount",
        "withdrawal",
        "withdrawals",
        "withdrawal amt",
        "withdrawal amount",
        "dr amount",
    ],
    "credit": [
        "credit",
        "credit amount",
        "deposit",
        "deposits",
        "deposit amt",
        "deposit amount",
        "cr amount",
    ],
    "balance": ["balance", "closing balance", "running balance", "balance amount"],
    "amount": ["amount", "transaction amount", "txn amount"],
    "type": ["type", "dr cr", "cr dr", "transaction type", "debit credit", "txn type"],
}


def money(value):
    if value is None or str(value).strip() in ("", "-", "--", "None", "nan"):
        return None
    text = (
        re.sub(r"(?i)INR|RS\.?|₹|\bCR\b|\bDR\b", "", str(value))
        .replace(",", "")
        .strip()
    )
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    try:
        amount = Decimal(text)
        if not amount.is_finite() or abs(amount) > Decimal("1000000000000"):
            raise ValueError("Invalid amount")
        return float(amount.quantize(Decimal(".01")))
    except InvalidOperation:
        raise ValueError("Invalid amount") from None


def date(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    text = str(value or "").strip()
    for fmt in (
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%d/%m/%y",
        "%d-%m-%Y",
        "%d-%m-%y",
        "%d %b %Y",
        "%d-%b-%Y",
        "%d-%b-%y",
        "%d.%m.%Y",
        "%Y-%m-%d %H:%M:%S",
    ):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    raise ValueError("Unrecognized transaction date")


def label(s):
    return re.sub(
        r"[^a-z0-9 ]", "", str(s or "").lower().replace("\n", " ").replace("/", " ")
    ).strip()


def table_records(rows, page=None):
    columns = {}
    start = None
    for i, row in enumerate(rows[:45]):
        labels = [label(x) for x in row]
        found = {
            k: next((j for j, x in enumerate(labels) if x in options), None)
            for k, options in ALIASES.items()
        }
        if (
            found["date"] is not None
            and found["description"] is not None
            and any(found[k] is not None for k in ("debit", "credit", "amount"))
        ):
            columns = found
            start = i + 1
            break
    if start is None:
        return [], []
    results = []
    warnings = []
    for index, row in enumerate(rows[start:], start + 1):

        def get(k):
            j = columns.get(k)
            return row[j] if j is not None and j < len(row) else None

        try:
            day = date(get("date"))
        except ValueError:
            continue
        try:
            debit = money(get("debit")) or 0
            credit = money(get("credit")) or 0
            if debit < 0 or credit < 0 or (debit and credit):
                raise ValueError("Ambiguous debit/credit columns")
            if not debit and not credit:
                amount = money(get("amount"))
                direction = label(get("type"))
                if amount is None or direction not in ("dr", "debit", "cr", "credit"):
                    raise ValueError("Missing explicit debit/credit direction")
                if direction in ("dr", "debit"):
                    debit = abs(amount)
                else:
                    credit = abs(amount)
            if not debit and not credit:
                continue
            narration = str(get("description") or "").strip()
            if not narration:
                raise ValueError("Missing narration")
            results.append(
                {
                    "txn_date": day,
                    "narration": narration,
                    "description": narration,
                    "amount": debit or credit,
                    "txn_type": "debit" if debit else "credit",
                    "balance": money(get("balance")),
                    "raw_debit": debit,
                    "raw_credit": credit,
                    "source_page": page,
                    "source_row": index,
                }
            )
        except ValueError as error:
            warnings.append(f"Row {index}: {error}. Review the original statement.")
    return results, warnings


def text_records(text, page):
    # Explicit DR/CR line format only. Never infer money direction from names.
    rows = []
    pattern = re.compile(
        r"^(\d{2}[/-]\d{2}[/-]\d{2,4}|\d{4}-\d{2}-\d{2})\s+(.+?)\s+([\d,]+\.\d{2})\s+(DR|CR|DEBIT|CREDIT)\s+([\d,]+\.\d{2})\s*$",
        re.I,
    )
    for i, line in enumerate(text.splitlines(), 1):
        m = pattern.match(line.strip())
        if m:
            day, description, amount, direction, balance = m.groups()
            rows.append(
                {
                    "txn_date": date(day),
                    "description": description,
                    "narration": description,
                    "amount": money(amount),
                    "txn_type": "debit"
                    if direction.upper() in ("DR", "DEBIT")
                    else "credit",
                    "balance": money(balance),
                    "source_page": page,
                    "source_row": i,
                }
            )
    return rows


def parse_statement(data, filename, password=""):
    suffix = Path(filename).suffix.lower()
    if not data or len(data) > 15_000_000:
        raise ValueError("Choose a statement under 15 MB.")
    records = []
    warnings = []
    texts = []
    parser = "tabular"
    ocr = False
    if suffix == ".pdf":
        import fitz
        import pdfplumber

        with fitz.open(stream=data, filetype="pdf") as document:
            if document.needs_pass and not document.authenticate(password):
                raise ValueError("The PDF password is missing or incorrect.")
            if len(document) > 80:
                raise ValueError("Split this PDF into files of at most 80 pages.")
            unlocked = document.tobytes(encryption=fitz.PDF_ENCRYPT_NONE)
            with pdfplumber.open(io.BytesIO(unlocked)) as pdf:
                for n, p in enumerate(document, 1):
                    text = p.get_text()
                    texts.append(text)
                    page_rows = []
                    for table in pdf.pages[n - 1].extract_tables():
                        found, notes = table_records(table, n)
                        page_rows.extend(found)
                        warnings.extend(notes)
                    if not page_rows:
                        page_rows = text_records(text, n)
                    if not text.strip():
                        if not shutil.which("tesseract"):
                            raise ValueError(
                                "This is a scanned PDF. Tesseract OCR must be installed on the server."
                            )
                        if len(document) > 20:
                            raise ValueError(
                                "Scanned PDFs are limited to 20 pages per upload."
                            )
                        from PIL import Image
                        import pytesseract

                        pix = p.get_pixmap(matrix=fitz.Matrix(2, 2))
                        image = Image.open(io.BytesIO(pix.tobytes("png")))
                        text = pytesseract.image_to_string(
                            image, config="--psm 6", timeout=15
                        )
                        texts[-1] = text
                        ocr = True
                        page_rows = text_records(text, n)
                    records.extend(page_rows)
            if not records:
                import hni_accounting_system as core

                text = "\n".join(texts)
                adapters = [
                    ("Axis", core._parse_axis_pdf_text),
                    ("AU", core._parse_au_pdf_text),
                    ("HDFC", core._parse_hdfc_pdf_text),
                ]
                for bank, fn in adapters:
                    if re.search(BANKS[bank], text, re.I):
                        result = fn(texts)
                        records = result.get("records", [])
                        parser = f"{bank} text adapter"
                        break
                if not records and re.search(BANKS["PNB"], text, re.I):
                    records = core._parse_pnb_pdf_text(text).get("records", [])
                    parser = "PNB text adapter"
                if not records and re.search(BANKS["Union"], text, re.I):
                    records = core._parse_union_bank_pdf_text(text).get("records", [])
                    parser = "Union text adapter"
        if ocr:
            warnings.append(
                "OCR was used. Verify every amount and narration against the original scan."
            )
    elif suffix in (".xlsx", ".xls"):
        import msoffcrypto

        try:
            office = msoffcrypto.OfficeFile(io.BytesIO(data))
            if office.is_encrypted():
                if not password:
                    raise ValueError("Enter the workbook password.")
                office.load_key(password=password)
                out = io.BytesIO()
                office.decrypt(out)
                data = out.getvalue()
        except ValueError:
            raise
        except Exception as error:
            if password:
                raise ValueError(
                    "The workbook could not be unlocked. Check its password and format."
                ) from error
        if zipfile.is_zipfile(io.BytesIO(data)):
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if (
                    sum(i.file_size for i in z.infolist()) > 80_000_000
                    or len(z.infolist()) > 2000
                ):
                    raise ValueError("Workbook expands beyond the processing limit.")
            import openpyxl

            wb = openpyxl.load_workbook(
                io.BytesIO(data), read_only=True, data_only=True
            )
            for sheet in wb:
                if sheet.max_row and sheet.max_row > 20000:
                    raise ValueError("A sheet exceeds the 20,000-row limit.")
                rows = list(sheet.values)
                texts.append(" ".join(str(x or "") for row in rows[:20] for x in row))
                found, notes = table_records(rows)
                records.extend(found)
                warnings.extend(notes)
            wb.close()
        else:
            import xlrd

            wb = xlrd.open_workbook(file_contents=data)
            for sheet in wb.sheets():
                if sheet.nrows > 20000:
                    raise ValueError("A sheet exceeds the 20,000-row limit.")
                rows = [
                    [
                        xlrd.xldate_as_datetime(cell.value, wb.datemode)
                        if cell.ctype == xlrd.XL_CELL_DATE
                        else cell.value
                        for cell in sheet.row(i)
                    ]
                    for i in range(sheet.nrows)
                ]
                texts.append(" ".join(str(x or "") for row in rows[:20] for x in row))
                found, notes = table_records(rows)
                records.extend(found)
                warnings.extend(notes)
    elif suffix == ".csv":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ValueError("Save the CSV as UTF-8.") from None
        rows = list(csv.reader(io.StringIO(text)))
        if len(rows) > 20001:
            raise ValueError("CSV files are limited to 20,000 rows.")
        texts = [text[:3000]]
        records, warnings = table_records(rows)
    else:
        raise ValueError("Upload a PDF, XLSX, XLS or CSV statement.")
    if not records:
        raise ValueError(
            "No unambiguous transactions were found. Use the CSV template or check the statement layout. No entries were posted."
        )
    if len(records) > 20000:
        raise ValueError("Split the statement into fewer than 20,000 transactions.")
    combined = "\n".join(texts) + " " + filename
    opening_source = "inferred from first transaction"
    closing_source = "last transaction balance"
    bank = next(
        (name for name, pattern in BANKS.items() if re.search(pattern, combined, re.I)),
        "Other bank",
    )
    records.sort(key=lambda r: r.get("txn_date", ""))
    first, last = records[0], records[-1]
    opening = first.get("balance")
    if opening is not None:
        opening = round(
            opening + first["amount"] * (1 if first["txn_type"] == "debit" else -1), 2
        )
    closing = last.get("balance")
    for field in ("opening", "closing"):
        match = re.search(
            r"(?i)"
            + field
            + r"\s+balance\s*[:=]?\s*(?:INR|Rs\.?|₹)?\s*([\d,]+\.\d{2})",
            combined,
        )
        if match:
            if field == "opening":
                opening = money(match.group(1))
                opening_source = "statement summary"
            else:
                closing = money(match.group(1))
                closing_source = "statement summary"
    if opening_source.startswith("inferred") and opening is not None:
        warnings.append(
            "Opening balance is inferred from the first row. Confirm it against the statement summary for independent reconciliation."
        )
    return {
        "opening_balance_source": opening_source,
        "closing_balance_source": closing_source,
        "records": records,
        "opening_balance": opening,
        "closing_balance": closing,
        "statement_from_date": first["txn_date"],
        "statement_to_date": last["txn_date"],
        "bank": bank,
        "parser": parser,
        "ocr": ocr,
        "warnings": warnings,
    }
