import io
import json
import time
from concurrent.futures import ThreadPoolExecutor
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from conftest import CSV, HEADERS, PASSWORD, post, register, upload
from goldilocks import accounting, service


def pending(c):
    return c.get("/api/pending").json["rows"]


def test_legacy_settlement_categories_cannot_bypass_accounting_review(client):
    upload(client, CSV)
    debit = next(r for r in pending(client) if r["txn_type"] == "debit")
    for key in ("exp_loan_emi", "exp_credit_card"):
        assert (
            post(
                client, "/api/rules", {"pattern": "SETTLEMENT", "ledger_key": key}
            ).status_code
            == 422
        )
        assert (
            post(
                client,
                "/api/transactions/" + debit["id"] + "/split",
                {
                    "allocations": [
                        {"key": key, "amount": 100},
                        {"key": "exp_loan_interest", "amount": 25.25},
                    ]
                },
            ).status_code
            == 422
        )
    assert client.get("/api/trace/" + debit["id"]).json["journal"] == []


def test_explicit_ephemeral_host_is_never_advertised_as_persistent(client, monkeypatch):
    monkeypatch.setenv("HNI_STORAGE_MODE", "ephemeral")
    monkeypatch.setenv("HNI_DATA_DIR", "/var/data")
    health = client.get("/api/health").json
    assert health["persistent_storage_configured"] is False
    assert "temporary storage" in health["storage_warning"]


def approved(c):
    return c.get("/api/ledger").json["rows"]


def test_account_isolation_and_no_global_login(app, client):
    r = upload(client, CSV)
    assert r.status_code == 201, r.json
    ids = [r["id"] for r in pending(client)]
    other, _ = register(app)
    assert len(pending(client)) == 2 and pending(other) == []
    assert other.get("/api/trace/" + ids[0]).status_code == 404
    assert (
        other.get("/api/statements/" + r.json["batch_id"] + "/source").status_code
        == 404
    )
    assert (
        post(
            other,
            "/api/transactions/" + ids[0] + "/classification",
            {"key": "exp_food"},
        ).status_code
        == 404
    )
    assert post(other, "/api/approve", {"ids": ids}).status_code == 404
    assert approved(client) == []
    assert len(pending(client)) == 2


def test_csrf_and_legacy_public_routes_closed(app, client):
    assert client.post("/api/sample").status_code == 403
    assert (
        client.post(
            "/api/sample", headers={**HEADERS, "Origin": "https://wrong.example"}
        ).status_code
        == 403
    )
    anonymous = app.test_client()
    assert anonymous.get("/api/pending").status_code == 401
    assert anonymous.get("/api/statements/nope/source").status_code == 401
    assert anonymous.post("/forgot-password", headers=HEADERS).status_code == 404
    assert anonymous.get("/download-upload/private.csv").status_code == 404
    assert anonymous.get("/users").status_code == 404
    assert "script-src 'self'" in client.get("/").headers["Content-Security-Policy"]
    assert "no-store" in client.get("/api/ledger").headers["Cache-Control"]
    assert client.get("/").status_code == 200


def test_password_recovery_rotates_code_and_revokes_sessions(app):
    c, data = register(app)
    second = app.test_client()
    assert (
        post(
            second, "/api/login", {"email": "other@example.com", "password": PASSWORD}
        ).status_code
        == 200
    )
    wrong = post(
        c,
        "/api/recover",
        {
            "email": "other@example.com",
            "recovery_code": "bad",
            "password": PASSWORD,
            "confirm_password": PASSWORD,
        },
    )
    assert wrong.status_code == 403
    reset = post(
        c,
        "/api/recover",
        {
            "email": "other@example.com",
            "recovery_code": data["recovery_code"],
            "password": "Another-fictional-pass",
            "confirm_password": "Another-fictional-pass",
        },
    )
    assert reset.status_code == 200, reset.json
    assert reset.json["recovery_code"] != data["recovery_code"]
    assert second.get("/api/pending").status_code == 401
    assert (
        post(
            c,
            "/api/recover",
            {
                "email": "other@example.com",
                "recovery_code": data["recovery_code"],
                "password": PASSWORD,
                "confirm_password": PASSWORD,
            },
        ).status_code
        == 403
    )
    assert post(c, "/api/logout", {}).status_code == 200
    assert c.get("/api/pending").status_code == 401
    assert (
        post(
            c, "/api/login", {"email": "other@example.com", "password": PASSWORD}
        ).status_code
        == 401
    )


def test_password_confirmation_and_secure_cookie(app):
    c = app.test_client()
    data = {
        "name": "Fictional Test",
        "email": "password@example.com",
        "password": PASSWORD,
        "confirm_password": "mismatch",
    }
    assert post(c, "/api/register", data).status_code == 422
    data["password"] = data["confirm_password"] = "short"
    assert post(c, "/api/register", data).status_code == 422
    app.config["COOKIE_SECURE"] = True
    data["password"] = data["confirm_password"] = PASSWORD
    r = post(c, "/api/register", data)
    assert r.status_code == 201
    cookie = r.headers["Set-Cookie"]
    assert all(v in cookie for v in ["Secure", "HttpOnly", "SameSite=Strict"])
    db = app.extensions["ledger"]["db"]
    with db._conn() as con:
        assert (
            con.execute("SELECT password_hash FROM users")
            .fetchone()[0]
            .startswith("scrypt:")
        )
        assert (
            len(con.execute("SELECT token_hash FROM web_sessions").fetchone()[0]) == 64
        )


def test_idle_timeout(app, client):
    db = app.extensions["ledger"]["db"]
    with db._conn() as c:
        c.execute("UPDATE web_sessions SET last_seen=?", (int(time.time()) - 1801,))
    assert client.get("/api/ledger").status_code == 401


def test_legacy_password_migration_accepts_old_short_password(app):
    import hni_accounting_system as core

    db = app.extensions["ledger"]["db"]
    with db._conn() as c:
        c.execute(
            "INSERT INTO users(id,name,email,phone,created_at,password_hash) VALUES(?,?,?,?,?,?)",
            (
                "legacy",
                "Legacy fixture",
                "legacy@example.com",
                "1234567890",
                "2026-01-01",
                core._hash_password("old-pass"),
            ),
        )
    c = app.test_client()
    r = post(c, "/api/login", {"email": "legacy@example.com", "password": "old-pass"})
    assert r.status_code == 200, r.json
    with db._conn() as dbcon:
        assert (
            dbcon.execute("SELECT password_hash FROM users WHERE id='legacy'")
            .fetchone()[0]
            .startswith("scrypt:")
        )


def test_atomic_approval_journal_trace_and_idempotency(client):
    r = upload(client, CSV, opening=1000, closing=10874.75)
    assert r.status_code == 201, r.json
    rows = pending(client)
    ids = [t["id"] for t in rows]
    assert approved(client) == []
    assert post(client, "/api/approve", {"ids": ids + ["unknown"]}).status_code == 404
    assert approved(client) == []
    assert post(client, "/api/approve", {"ids": ids}).json["approved"] == 2
    assert post(client, "/api/approve", {"ids": ids}).json["approved"] == 0
    assert len(approved(client)) == 2
    report = client.get("/api/reports").json
    assert report["income_total"] == 10000 and report["expense_total"] == 125.25
    assert report["net_worth"] == 10874.75 and report["difference"] == 0
    assert report["journal_balanced"] and report["missing_journals"] == 0
    for rid in ids:
        trace = client.get("/api/trace/" + rid).json
        assert len(trace["journal"]) == 2
        assert sum(x["debit"] for x in trace["journal"]) == sum(
            x["credit"] for x in trace["journal"]
        )
        assert trace["source"]["raw_description"]
        assert trace["events"][0]["action"] == "transaction_posted"
    assert (
        client.get("/api/statements/" + r.json["batch_id"] + "/source").data
        == CSV.encode()
    )


def test_new_import_does_not_clear_pending_and_duplicate_rules(client):
    assert upload(client, CSV).json["staged"] == 2
    assert upload(client, CSV).json["duplicate"] is True
    second = CSV.replace("2026-01-01", "2026-02-01").replace("2026-01-02", "2026-02-02")
    assert upload(client, second, name="feb.csv").json["staged"] == 2
    assert len(pending(client)) == 4
    overlap = CSV + "2026-01-03,SWIGGY FOOD ORDER,25.00,,10849.75\n"
    r = upload(client, overlap, name="overlap.csv").json
    assert r["staged"] == 1 and r["duplicates"] == 2
    assert len(pending(client)) == 5


def test_legitimate_identical_payments_are_not_collapsed(client):
    text = "Date,Narration,Debit,Credit,Balance\n2026-01-01,SWIGGY FOOD ORDER,50,,950\n2026-01-01,SWIGGY FOOD ORDER,50,,900\n"
    assert upload(client, text).json["staged"] == 2
    assert len(pending(client)) == 2


def test_concurrent_duplicate_import_is_atomic(app, client):
    uid = client.get("/api/session").json["user"]["id"]
    db = app.extensions["ledger"]["db"]
    edb = app.extensions["ledger"]["edb"]
    from goldilocks.ingestion import parse_statement

    parsed = parse_statement(CSV.encode(), "test.csv")

    def run(_):
        return service.stage(
            db, edb, uid, "Concurrent", "test.csv", CSV.encode(), parsed
        )

    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(run, range(2)))
    assert sum(x["staged"] for x in results) == 2
    assert len(pending(client)) == 2


def test_personal_rules_and_learning_are_private(app, client):
    r = post(
        client,
        "/api/rules",
        {
            "pattern": "CEDAR WEEKLY",
            "ledger_key": "exp_food",
            "txn_type": "debit",
            "match_mode": "contains",
        },
    )
    assert r.status_code == 201, r.json
    text = "Date,Narration,Debit,Credit,Balance\n2026-01-01,CEDAR WEEKLY XYZ,100,,900\n"
    upload(client, text)
    row = pending(client)[0]
    assert (
        row["predicted_ledger_key"] == "exp_food" and row["source"] == "personal rule"
    )
    post(
        client,
        "/api/transactions/" + row["id"] + "/classification",
        {"key": "exp_grocery", "note": "Fixture correction"},
    )
    upload(client, text.replace("2026-01-01", "2026-01-02"), name="next.csv")
    new = next(r for r in pending(client) if r["txn_date"] == "2026-01-02")
    assert (
        new["predicted_ledger_key"] == "exp_grocery"
        and new["source"] == "reviewed example"
    )
    other, _ = register(app)
    assert other.get("/api/learning-examples").json == []
    assert other.get("/api/rules").json["rows"] == []
    assert (
        other.delete("/api/rules/" + r.json["row"]["id"], headers=HEADERS).status_code
        == 404
    )


def test_unclassified_cannot_post_and_splits_balance(client):
    text = "Date,Narration,Debit,Credit,Balance\n2026-01-01,HOME LOAN EMI,1000,,9000\n"
    upload(client, text, opening=10000, closing=9000)
    row = pending(client)[0]
    assert row["section"] == "Suspense"
    assert post(client, "/api/approve", {"ids": [row["id"]]}).status_code == 422
    bad = {
        "allocations": [
            {"key": "liability_loan_outstanding", "amount": 700},
            {"key": "exp_loan_interest", "amount": 200},
        ]
    }
    assert (
        post(client, "/api/transactions/" + row["id"] + "/split", bad).status_code
        == 422
    )
    assert (
        client.put(
            "/api/openings",
            json={
                "account_id": "Test account",
                "key": "liability_loan_outstanding",
                "amount": 5000,
                "as_of": "2026-01-01",
            },
            headers=HEADERS,
        ).status_code
        == 200
    )
    good = {
        "allocations": [
            {"key": "liability_loan_outstanding", "amount": 800},
            {"key": "exp_loan_interest", "amount": 200},
        ]
    }
    assert (
        post(client, "/api/transactions/" + row["id"] + "/split", good).status_code
        == 200
    )
    assert post(client, "/api/approve", {"ids": [row["id"]]}).json["approved"] == 1
    r = client.get("/api/reports").json
    assert (
        r["expense_total"] == 200
        and r["totals"]["Liabilities"] == 4200
        and r["net_worth"] == 4800
        and r["difference"] == 0
    )
    assert len(client.get("/api/trace/" + row["id"]).json["journal"]) == 3


def test_posted_correction_rebuilds_journal_keeps_source(client):
    upload(client, CSV)
    rows = pending(client)
    post(client, "/api/approve", {"ids": [r["id"] for r in rows]})
    debit = next(r for r in rows if r["txn_type"] == "debit")
    rid = debit["id"]
    original = client.get("/api/trace/" + rid).json["source"]
    r = post(
        client,
        "/api/transactions/" + rid + "/classification",
        {
            "key": "asset_investment_mf",
            "posted": True,
            "note": "Fixture capital allocation",
        },
    )
    assert r.status_code == 200, r.json
    data = client.get("/api/trace/" + rid).json
    assert data["source"] == original
    assert len(data["journal"]) == 2 and any(
        j["code"] == "asset_investment_mf" for j in data["journal"]
    )
    assert data["events"][-1]["action"] == "posted_classification_corrected"
    assert client.get("/api/reports").json["expense_total"] == 0
    assert (
        post(
            client,
            "/api/transactions/" + rid + "/classification",
            {"key": "suspense_debit", "posted": True},
        ).status_code
        == 422
    )


def test_reconciliation_distinguishes_inferred_and_wrong_balances(client):
    r = upload(client, CSV)
    recon = client.get("/api/reconciliation").json["rows"][0]
    assert recon["status"] == "consistent_inferred"
    rid = r.json["batch_id"]
    assert (
        client.put(
            "/api/statements/" + rid + "/balances",
            json={"opening": 1000, "closing": 10000},
            headers=HEADERS,
        ).status_code
        == 200
    )
    recon = client.get("/api/reconciliation").json["rows"][0]
    assert recon["status"] == "gap" and recon["difference"] == 874.75
    client.put(
        "/api/statements/" + rid + "/balances",
        json={"opening": 1000, "closing": 10874.75},
        headers=HEADERS,
    )
    assert client.get("/api/reconciliation").json["rows"][0]["status"] == "matched"
    broken = CSV.replace("10874.75", "10873.75")
    upload(client, broken, account="Other account", opening=1000, closing=10874.75)
    recon = next(
        r
        for r in client.get("/api/reconciliation").json["rows"]
        if r["account_id"] == "Other account"
    )
    assert recon["status"] == "gap" and recon["row_breaks"] == [2]


def test_date_account_filters_and_exports(client):
    upload(client, CSV)
    post(client, "/api/approve", {"ids": [r["id"] for r in pending(client)]})
    period = client.get("/api/reports?from=2026-01-02&to=2026-01-02").json
    assert (
        period["income_total"] == 0
        and period["expense_total"] == 125.25
        and period["net_worth"] == 10874.75
    )
    assert client.get("/api/reports?account=Unknown").json["income_total"] == 0
    assert len(client.get("/api/ledger?from=2026-01-02").json["rows"]) == 1
    for kind in ["csv", "xlsx", "pdf"]:
        r = client.get("/api/export/" + kind + "?from=2026-01-02")
        assert r.status_code == 200 and len(r.data) > 50
        if kind == "csv":
            assert "SALARY" not in r.data.decode("utf-8-sig")
        if kind == "xlsx":
            assert load_workbook(io.BytesIO(r.data)).active.max_row == 2
        if kind == "pdf":
            assert "Goldilocks" in PdfReader(io.BytesIO(r.data)).pages[0].extract_text()
    assert client.get("/api/reports?from=not-a-date").status_code == 422
    assert client.get("/api/reports?from=2026-02-01&to=2026-01-01").status_code == 422


def test_exports_neutralize_spreadsheet_formulas(client):
    text = "Date,Narration,Debit,Credit,Balance\n2026-01-01,=1+1,100,,900\n"
    upload(client, text)
    row = pending(client)[0]
    post(
        client,
        "/api/transactions/" + row["id"] + "/classification",
        {"key": "exp_food"},
    )
    post(client, "/api/approve", {"ids": [row["id"]]})
    assert "'=1+1" in client.get("/api/export/csv").data.decode("utf-8-sig")
    ws = load_workbook(io.BytesIO(client.get("/api/export/xlsx").data)).active
    assert ws.cell(2, 3).data_type == "s"


def test_model_bootstraps_and_obvious_rules(client):
    meta = client.get("/api/meta").json
    assert meta["model"]["ready"] and meta["model"]["classes"] > 20
    sample = post(client, "/api/sample", {})
    assert sample.status_code == 201, sample.json
    rows = pending(client)
    assert len(rows) == 80 and all(r["section"] != "Suspense" for r in rows)
    assert (
        next(r for r in rows if "BANK INTEREST" in r["narration"])[
            "predicted_ledger_key"
        ]
        == "income_interest"
    )
    assert (
        next(r for r in rows if "RENT" in r["narration"])["predicted_ledger_key"]
        == "exp_rent"
    )


@pytest.mark.parametrize(
    "value",
    [[], ["same", "same"], None, "one", {"id": 1}, [1], list(map(str, range(1001)))],
)
def test_invalid_approval_selections_are_rejected(client, value):
    assert post(client, "/api/approve", {"ids": value}).status_code == 422


def test_login_rate_limit_and_generic_errors(app):
    c = app.test_client()
    for _ in range(15):
        assert (
            post(
                c, "/api/login", {"email": "missing@example.com", "password": "wrong"}
            ).status_code
            == 401
        )
    r = post(c, "/api/login", {"email": "missing@example.com", "password": "wrong"})
    assert r.status_code == 429
    assert "traceback" not in json.dumps(r.json).lower()


def test_running_amount_equal_to_balance_is_valid(client):
    text = "Date,Narration,Debit,Credit,Balance\n2026-01-01,SALARY CEDAR,,1000,1000\n"
    assert upload(client, text, opening=0, closing=1000).json["staged"] == 1


def test_integer_paise_financial_identity():
    assert accounting.paise("0.29") == 29 and accounting.paise("100.005") == 10001
    for value in ["NaN", "Infinity", "1000000000001"]:
        with pytest.raises(ValueError):
            accounting.paise(value)


def test_identical_posted_rows_survive_application_restart(app, client):
    text = "Date,Narration,Debit,Credit,Balance\n2026-01-01,SWIGGY FOOD ORDER,50,,950\n2026-01-01,SWIGGY FOOD ORDER,50,,900\n"
    r = upload(client, text)
    assert r.status_code == 201, r.json
    ids = [r["id"] for r in pending(client)]
    assert post(client, "/api/approve", {"ids": ids}).json["approved"] == 2
    from goldilocks.app import create_app

    reopened = create_app(
        app.extensions["ledger"]["db"].db_path, testing=True
    ).test_client()
    assert (
        post(
            reopened,
            "/api/login",
            {"email": "tester@example.com", "password": PASSWORD},
        ).status_code
        == 200
    )
    assert len(approved(reopened)) == 2
    assert reopened.get("/api/reports").json["expense_total"] == 100
